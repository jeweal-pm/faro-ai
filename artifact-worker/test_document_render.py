import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from document_render import render_document, validate_input


class DocumentRenderTests(unittest.TestCase):
    def test_formats_preserve_input_and_prevent_excel_formula_execution(self):
        sample = {"title": "รายงาน Test", "columns": ["Product", "Price"],
                  "rows": [["=HYPERLINK(\"https://invalid.test\")", 125], ["แหวน", 0]]}
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for fmt in ("pdf", "xlsx", "docx", "pptx"):
                path = root / ("report." + fmt)
                result = render_document(fmt, sample, path)
                self.assertEqual(result["row_count"], 2)
                self.assertGreater(path.stat().st_size, 0)
            workbook = load_workbook(root / "report.xlsx", data_only=False)
            self.assertEqual(workbook.active["A3"].data_type, "s")
            self.assertTrue(workbook.active["A3"].value.startswith("'="))
            self.assertEqual(workbook.active["B3"].value, 125)
            self.assertEqual(Document(root / "report.docx").tables[0].rows[2].cells[0].text, "แหวน")
            self.assertEqual(len(Presentation(root / "report.pptx").slides), 1)

    def test_rejects_malformed_and_unbounded_inputs(self):
        cases = [
            {"title": "x", "columns": ["a"], "rows": [[1, 2]]},
            {"title": "x", "columns": ["a"], "rows": [[float("nan")]]},
            {"title": "x", "columns": ["a"], "rows": [[{}]]},
            {"title": "x", "columns": ["a"], "rows": [["x" * 501]]},
            {"title": "x", "columns": ["a"], "rows": [], "path": "other"},
        ]
        for case in cases:
            with self.subTest(case=case), self.assertRaises(ValueError):
                validate_input(case)

    def test_presentation_refuses_oversized_cells_without_truncating(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "slide_cell_too_long"):
                render_document("pptx", {"title": "x", "columns": ["a"],
                                         "rows": [["x" * 91]]}, Path(folder) / "x.pptx")


if __name__ == "__main__":
    unittest.main()
