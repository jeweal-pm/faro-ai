"""Fixed, bounded document renderers. Input is data; it is never executable code."""
from __future__ import annotations

import math
import os
import re
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


FORMATS = {"pdf", "xlsx", "docx", "pptx"}
MAX_ROWS, MAX_COLUMNS, MAX_TEXT = 500, 20, 500
MAX_BYTES = 10 * 1024 * 1024


def _text(value: Any, limit: int = MAX_TEXT) -> str:
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, bool)):
        raise ValueError("unsupported_cell_type")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("invalid_number")
    result = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(value))
    if len(result) > limit:
        raise ValueError("cell_too_long")
    return result


def validate_input(document: dict[str, Any]) -> tuple[str, list[str], list[list[Any]]]:
    if not isinstance(document, dict) or set(document) != {"title", "columns", "rows"}:
        raise ValueError("invalid_document")
    title = _text(document["title"], 120).strip()
    cols, rows = document["columns"], document["rows"]
    if not title or not isinstance(cols, list) or not 1 <= len(cols) <= MAX_COLUMNS:
        raise ValueError("invalid_columns")
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        raise ValueError("invalid_rows")
    columns = [_text(col, 120).strip() for col in cols]
    if any(not col for col in columns):
        raise ValueError("invalid_columns")
    checked: list[list[Any]] = []
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise ValueError("invalid_row_width")
        for cell in row:
            _text(cell)
        checked.append(row)
    return title, columns, checked


def spreadsheet_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        _text(value)
        return value
    text = _text(value)
    # Excel treats a formula after whitespace/control characters as executable.
    return "'" + text if re.match(r"^[\s]*[=+@-]", text) else text


def _font_path() -> Path:
    candidates = [os.environ.get("FAROAI_DOCUMENT_FONT", ""),
                  "/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf",
                  "C:/Windows/Fonts/LeelawUI.ttf"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise ValueError("thai_font_unavailable")


def _pdf(path: Path, title: str, columns: list[str], rows: list[list[Any]]) -> None:
    if "FaroDocument" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("FaroDocument", str(_font_path())))
    page = landscape(A3) if len(columns) > 6 else A4
    style = ParagraphStyle("FaroCell", fontName="FaroDocument", fontSize=7.5,
                           leading=11, wordWrap="CJK", alignment=TA_LEFT)
    heading = ParagraphStyle("FaroHeading", fontName="FaroDocument", fontSize=16,
                             leading=22, wordWrap="CJK")
    table_rows = [[Paragraph(escape(c), style) for c in columns]]
    table_rows += [[Paragraph(escape(_text(cell)).replace("\n", "<br/>"), style)
                    for cell in row] for row in rows]
    width = (page[0] - 72) / len(columns)
    table = LongTable(table_rows, colWidths=[width] * len(columns), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dff2f1")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cad8e3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    SimpleDocTemplate(str(path), pagesize=page, leftMargin=36, rightMargin=36,
                      topMargin=36, bottomMargin=36).build([
                          Paragraph(escape(title), heading), Spacer(1, 14), table])


def _xlsx(path: Path, title: str, columns: list[str], rows: list[list[Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([spreadsheet_value(title)])
    sheet.append([spreadsheet_value(value) for value in columns])
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = PatternFill("solid", fgColor="DFF2F1")
    for row in rows:
        sheet.append([spreadsheet_value(value) for value in row])
    sheet.freeze_panes = "A3"
    workbook.save(path)


def _docx(path: Path, title: str, columns: list[str], rows: list[list[Any]]) -> None:
    document = Document()
    document.add_heading(title, 0)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, columns, strict=True):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = _text(value)
    document.save(path)


def _pptx(path: Path, title: str, columns: list[str], rows: list[list[Any]]) -> None:
    # A fixed page budget avoids silently clipping long data or generating huge decks.
    if any(len(value) > 50 for value in columns) or any(len(_text(value)) > 90 for row in rows for value in row):
        raise ValueError("slide_cell_too_long")
    groups = [columns[i:i + 5] for i in range(0, len(columns), 5)]
    row_pages = [rows[i:i + 8] for i in range(0, len(rows), 8)] or [[]]
    if len(groups) * len(row_pages) > 64:
        raise ValueError("slide_limit_exceeded")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for group_index, group in enumerate(groups):
        for page_index, page in enumerate(row_pages):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(.5), Inches(.25), Inches(12.3), Inches(.65))
            paragraph = heading.text_frame.paragraphs[0]
            paragraph.text = title
            paragraph.font.size = Pt(24)
            table = slide.shapes.add_table(len(page) + 1, len(group), Inches(.5), Inches(1.15),
                                           Inches(12.3), Inches(.48 * (len(page) + 1))).table
            for index, value in enumerate(group):
                table.cell(0, index).text = value
            for row_index, row in enumerate(page, 1):
                for col_index in range(len(group)):
                    table.cell(row_index, col_index).text = _text(row[group_index * 5 + col_index])
            for table_row in table.rows:
                for cell in table_row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.size = Pt(12)
    presentation.save(path)


def render_document(kind: str, document: dict[str, Any], path: Path) -> dict[str, Any]:
    if kind not in FORMATS:
        raise ValueError("unsupported_format")
    title, columns, rows = validate_input(document)
    {"pdf": _pdf, "xlsx": _xlsx, "docx": _docx, "pptx": _pptx}[kind](path, title, columns, rows)
    size = path.stat().st_size
    if size <= 0 or size > MAX_BYTES:
        raise ValueError("artifact_size_invalid")
    return {"row_count": len(rows), "column_count": len(columns), "size_bytes": size}
